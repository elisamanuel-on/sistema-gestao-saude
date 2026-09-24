"""Testes do Sistema de Gestão de Utentes e Cuidados de Saúde.

Cobre: limpeza/estruturação dos dados de diabetes (usada para treinar o
modelo), a escala clínica de risco de queda (Morse Fall Scale), o resumo
automático (motor de regras), o roteamento/callbacks principais da app
(lista de utentes, ficha, avaliação de risco), e os módulos do tipo de
unidade Clínica/Hospital (exames, plano de cuidados multidisciplinar,
consultas, faturação com seguros).

Não faz suposições sobre nomes concretos de utentes: usa sempre o primeiro
utente carregado a partir de dados/utentes.csv (sintético — ver
scripts/gerar_dados_utentes.py), para não ficar frágil a alterações na
geração dos dados.
"""
from io import BytesIO

import dash
import pandas as pd

import app as m
import constantes
import riscos
from limpeza import limpar_registo_bruto, preparar_para_modelo


def _primeiro_utente():
    return m.DADOS["utentes"].iloc[0]


def _primeiro_utente_com(chave):
    """Primeiro utente cuja tabela `chave` (em m.DADOS) tem pelo menos uma
    linha — para testar abas que dependem de haver dados (exames, plano)."""
    ids_com_dados = set(m.DADOS[chave]["id_utente"])
    return m.DADOS["utentes"][m.DADOS["utentes"]["id_utente"].isin(ids_com_dados)].iloc[0]


class FalsoNClicks:
    pass


# --- Limpeza / estruturação (dados de diabetes) -----------------------------


def test_limpeza_remove_duplicados_e_marca_invalidos():
    df, relatorio = limpar_registo_bruto("dados/registo_bruto.csv")
    assert relatorio["duplicados_removidos"] == 15
    assert relatorio["valores_invalidos_corrigidos"] == 6
    assert relatorio["linhas_estruturadas"] == 768


def test_limpeza_mantem_valores_em_falta_como_nan_nao_inventa():
    df, _ = limpar_registo_bruto("dados/registo_bruto.csv")
    assert df["insulina"].isna().sum() > 0


def test_preparar_para_modelo_nao_deixa_nan():
    df, _ = limpar_registo_bruto("dados/registo_bruto.csv")
    X, y = preparar_para_modelo(df)
    assert X.isna().sum().sum() == 0
    assert set(y.unique()) <= {0, 1}


def test_data_em_3_formatos_diferentes_e_reconhecida():
    bruto = pd.DataFrame(
        {
            "id_paciente": ["P1", "P2", "P3"],
            "data_registo": ["2024-01-15", "15-01-2024", "15/01/2024"],
            "gravidezes": ["1", "1", "1"],
            "glicose": ["120", "120", "120"],
            "pressao_arterial": ["70", "70", "70"],
            "espessura_pele": ["20", "20", "20"],
            "insulina": ["80", "80", "80"],
            "imc": ["25.0", "25.0", "25.0"],
            "pedigree_diabetes": ["0.3", "0.3", "0.3"],
            "idade": ["40", "40", "40"],
            "resultado": ["Sim", "sim", "SIM"],
        }
    )
    limpo, relatorio = limpar_registo_bruto(bruto)
    assert relatorio["datas_nao_reconhecidas"] == 0
    assert limpo["data_registo"].nunique() == 1  # as 3 datas são o mesmo dia


# --- Risco de queda (Morse Fall Scale) ---------------------------------------


def test_risco_queda_pontuacao_minima_e_baixo():
    pontuacao, classificacao, tom = riscos.calcular_risco_queda(
        "Não", "Não", "Nenhum / repouso / acompanhado por profissional", "Não",
        "Normal / acamado / cadeira de rodas", "Orientado(a), consciente das suas limitações",
    )
    assert pontuacao == 0
    assert classificacao == "Risco baixo"
    assert tom == "risco_baixo"


def test_risco_queda_pontuacao_maxima_e_elevado():
    pontuacao, classificacao, tom = riscos.calcular_risco_queda(
        "Sim, nos últimos 3 meses", "Sim (mais de um diagnóstico ativo)", "Apoia-se em mobiliário", "Sim",
        "Comprometida", "Sobrestima capacidades ou esquece limitações",
    )
    assert pontuacao == 125
    assert classificacao == "Risco elevado"
    assert tom == "risco_elevado"


def test_risco_queda_limiar_baixo_para_moderado():
    # histórico de quedas sozinho = 25 pontos -> fronteira exata baixo/moderado
    pontuacao, classificacao, _tom = riscos.calcular_risco_queda(
        "Sim, nos últimos 3 meses", "Não", "Nenhum / repouso / acompanhado por profissional", "Não",
        "Normal / acamado / cadeira de rodas", "Orientado(a), consciente das suas limitações",
    )
    assert pontuacao == 25
    assert classificacao == "Risco moderado"


def test_risco_queda_limiar_moderado_para_elevado():
    # soro (20) + marcha comprometida (20) + diagnóstico secundário (15) = 55 -> elevado
    pontuacao, classificacao, _tom = riscos.calcular_risco_queda(
        "Não", "Sim (mais de um diagnóstico ativo)", "Nenhum / repouso / acompanhado por profissional", "Sim",
        "Comprometida", "Orientado(a), consciente das suas limitações",
    )
    assert pontuacao == 55
    assert classificacao == "Risco elevado"


def test_classificar_risco_generico_limiares():
    assert riscos.classificar_risco(0.10)[0] == "Risco baixo"
    assert riscos.classificar_risco(0.50)[0] == "Risco moderado"
    assert riscos.classificar_risco(0.90)[0] == "Risco elevado"


# --- Resumo automático (motor de regras) -------------------------------------


def test_resumo_ia_sem_nenhum_resultado_nao_rebenta():
    resumo = riscos.gerar_resumo_ia("Ana Teste")
    assert "Ana Teste" in resumo
    assert "apoio à decisão" in resumo


def test_resumo_ia_marca_risco_elevado_para_revisao_prioritaria():
    resumo = riscos.gerar_resumo_ia(
        "Ana Teste",
        resultado_diabetes=(0.9, "Risco elevado", "glicose"),
    )
    assert "revisão clínica prioritária" in resumo
    assert "diabetes" in resumo


def test_resumo_ia_nunca_repete_a_palavra_risco_duas_vezes_seguidas():
    resumo = riscos.gerar_resumo_ia(
        "Ana Teste",
        resultado_diabetes=(0.9, "Risco elevado", "glicose"),
        resultado_cardio=(0.5, "Risco moderado", "idade"),
        resultado_queda=(10, "Risco baixo", "risco_baixo"),
    )
    assert "risco risco" not in resumo.lower()


def test_principais_fatores_ordena_por_peso_absoluto():
    importancias = {"a": 0.1, "b": -0.9, "c": 0.5}
    rotulos = {"a": "Fator A", "b": "Fator B", "c": "Fator C"}
    resultado = riscos.principais_fatores(importancias, rotulos, n=2)
    assert resultado == "Fator B, Fator C"


# --- App: dados carregados e páginas -----------------------------------------


def test_dados_utentes_carregados_tem_pelo_menos_um_utente():
    assert len(m.DADOS["utentes"]) > 0


def test_cada_utente_tem_avaliacao_diabetes_cardio_e_queda():
    utente = _primeiro_utente()
    assert utente["id_utente"] in m.DADOS["avaliacao_diabetes"].index
    assert utente["id_utente"] in m.DADOS["avaliacao_cardio"].index
    assert utente["id_utente"] in m.DADOS["avaliacao_queda"].index


def test_pagina_utentes_renderiza():
    pagina = m._pagina_utentes()
    assert pagina is not None


def test_pagina_ficha_utente_existente_mostra_nome():
    utente = _primeiro_utente()
    pagina = m._pagina_ficha_utente(utente["id_utente"])
    # o H1 com o nome é o segundo filho (depois da ligação "voltar")
    assert utente["nome"] in str(pagina)


def test_pagina_ficha_utente_inexistente_nao_rebenta():
    pagina = m._pagina_ficha_utente("ID-QUE-NAO-EXISTE")
    assert "não encontrado" in str(pagina).lower()


def test_filtrar_utentes_por_texto():
    utente = _primeiro_utente()
    primeiro_nome = utente["nome"].split()[0]
    resultado = m._filtrar_utentes(primeiro_nome, None, None)
    assert utente["nome"] in str(resultado)


def test_filtrar_utentes_sem_resultado():
    resultado = m._filtrar_utentes("Nome Que Nao Existe De Certeza Absoluta", None, None)
    assert "Nenhum utente encontrado" in str(resultado)


def test_formatar_euros_usa_formato_portugues():
    assert m._formatar_euros(1234.5, 2) == "1.234,50 €"
    assert m._formatar_euros(18755) == "18.755 €"


def test_pagina_ocupacao_renderiza():
    assert m._pagina_ocupacao() is not None


def test_pagina_faturacao_renderiza():
    assert m._pagina_faturacao() is not None


# --- Modelos de risco: previsões plausíveis -----------------------------------


def test_modelo_diabetes_prevé_probabilidade_valida():
    linha = pd.DataFrame([{c: 1.0 for c in m.RES_DIABETES}])
    prob = m.MODELO_DIABETES.predict_proba(linha)[0, 1]
    assert 0.0 <= prob <= 1.0


def test_modelo_cardio_prevé_probabilidade_valida():
    from scripts.treinar_modelo_cardio import COLUNAS_CARDIO

    linha = pd.DataFrame([{c: 1.0 for c in COLUNAS_CARDIO}])
    prob = m.MODELO_CARDIO.predict_proba(linha)[0, 1]
    assert 0.0 <= prob <= 1.0


def test_metricas_dos_modelos_tem_auc_razoavel():
    assert m.METRICAS_DIABETES["auc_conjunto_teste"] > 0.7
    assert m.METRICAS_CARDIO["auc_conjunto_teste"] > 0.7


# --- Módulo Clínica/Hospital: tipos, exames, plano de cuidados, consultas ----


def test_tipos_cuidado_inclui_clinica_hospital():
    assert "Clínica/Hospital" in constantes.TIPOS_CUIDADO


def test_slug_remove_acentos_espacos_e_barras():
    assert m._slug("Clínica/Hospital") == "clinica-hospital"
    assert m._slug("Concluído") == "concluido"
    assert m._slug("Em espera") == "em-espera"


def test_todos_os_utentes_tem_tipo_cuidado_valido():
    tipos_usados = set(m.DADOS["utentes"]["tipo_cuidado"].unique())
    assert tipos_usados <= set(constantes.TIPOS_CUIDADO)


def test_estado_internado_so_existe_em_ucc_erpi():
    internados = m.DADOS["utentes"][m.DADOS["utentes"]["estado"] == "Internado"]
    assert set(internados["tipo_cuidado"].unique()) <= {"UCC", "ERPI"}


def test_plano_cuidados_profissional_e_coerente_com_a_area():
    plano = m.DADOS["plano_cuidados"]
    for area, profissionais_da_area in plano.groupby("area_profissional")["profissional_responsavel"]:
        esperado = set(constantes.PROFISSIONAIS_POR_AREA_CUIDADOS[area])
        assert set(profissionais_da_area.unique()) <= esperado


def test_consulta_especialidade_bate_certo_com_o_profissional():
    consultas = m.DADOS["consultas"]
    for _idx, c in consultas.iterrows():
        assert constantes.PROFISSIONAL_ESPECIALIDADE[c["profissional"]] == c["especialidade"]


def test_aba_exames_renderiza_para_utente_com_exames():
    utente = _primeiro_utente_com("exames")
    assert m._aba_exames(utente["id_utente"]) is not None


def test_aba_plano_cuidados_renderiza_para_utente_com_plano():
    utente = _primeiro_utente_com("plano_cuidados")
    assert m._aba_plano_cuidados(utente["id_utente"]) is not None


def test_pagina_consultas_renderiza():
    assert m._pagina_consultas() is not None


def test_filtrar_consultas_por_especialidade():
    especialidade = m.DADOS["consultas"]["especialidade"].iloc[0]
    resultado = m._filtrar_consultas(None, especialidade, None)
    assert especialidade in str(resultado)


def test_filtrar_consultas_sem_resultado():
    resultado = m._filtrar_consultas("Nome Que Nao Existe De Certeza Absoluta", None, None)
    assert "Nenhuma consulta encontrada" in str(resultado)


def test_faturacao_tem_colunas_de_seguro():
    fat = m.DADOS["faturacao"]
    for coluna in ["seguradora", "numero_apolice", "cobertura_percentual_seguro", "valor_seguradora", "copagamento_utente"]:
        assert coluna in fat.columns


def test_faturacao_sem_seguro_aparece_como_travessao():
    fat = m.DADOS["faturacao"]
    # depois do fillna em _carregar_todos_os_dados, nunca deve sobrar NaN
    assert fat["seguradora"].isna().sum() == 0


# --- Visão Geral, login por perfil, PDF, calendário, histórico, sobre -------


def test_perfis_de_acesso_tem_quatro_opcoes():
    assert constantes.PERFIS_ACESSO == ["Enfermeiro", "Médico", "Receção", "Admin"]


def test_so_admin_ve_a_seccao_profissionais():
    assert "Profissionais" in constantes.SECCOES_POR_PERFIL["Admin"]
    assert "Profissionais" not in constantes.SECCOES_POR_PERFIL["Receção"]
    assert "Profissionais" not in constantes.SECCOES_POR_PERFIL["Enfermeiro"]
    assert "Profissionais" not in constantes.SECCOES_POR_PERFIL["Médico"]


def test_recepcao_e_admin_veem_faturacao_mas_nao_enfermeiro_nem_medico():
    assert "Faturação" in constantes.SECCOES_POR_PERFIL["Receção"]
    assert "Faturação" in constantes.SECCOES_POR_PERFIL["Admin"]
    assert "Faturação" not in constantes.SECCOES_POR_PERFIL["Enfermeiro"]
    assert "Faturação" not in constantes.SECCOES_POR_PERFIL["Médico"]


def test_pagina_login_tem_o_contentor_do_corpo_dinamico():
    # A partir de agora _pagina_login() só monta a casca do cartão — o
    # conteúdo (lista de perfis ou lista de profissionais) é preenchido por
    # um callback a partir de perfil-provisorio-login, para suportar o passo
    # extra de escolher qual profissional (Médico/Enfermeiro).
    pagina = str(m._pagina_login())
    assert "corpo-login" in pagina
    assert "perfil-provisorio-login" in pagina


def test_corpo_login_escolha_perfil_mostra_os_quatro_perfis():
    corpo = str(m._corpo_login_escolha_perfil())
    for perfil in constantes.PERFIS_ACESSO:
        assert perfil in corpo


def test_corpo_login_escolha_profissional_lista_medicos_ativos():
    corpo = str(m._corpo_login_escolha_profissional("Médico"))
    medicos_ativos = m.DADOS["profissionais"]
    medicos_ativos = medicos_ativos[(medicos_ativos["categoria"] == "Médico") & (medicos_ativos["estado"] == "Ativo")]
    assert not medicos_ativos.empty
    for nome in medicos_ativos["nome"]:
        assert nome in corpo


def test_dados_sessao_aceita_formato_novo_e_antigo():
    assert m._dados_sessao(None) == (None, None)
    assert m._dados_sessao("Médico") == ("Médico", None)  # sessões antigas (só string)
    assert m._dados_sessao({"perfil": "Médico", "profissional": "Dr. Hugo Teixeira"}) == ("Médico", "Dr. Hugo Teixeira")
    assert m._dados_sessao({"perfil": "Receção", "profissional": None}) == ("Receção", None)


def test_profissoes_com_profissional_sao_medico_e_enfermeiro():
    assert m.PERFIS_COM_PROFISSIONAL == {"Médico", "Enfermeiro"}
    assert "Receção" not in m.PERFIS_COM_PROFISSIONAL
    assert "Admin" not in m.PERFIS_COM_PROFISSIONAL


def test_renderizar_corpo_login_sem_provisorio_mostra_perfis():
    corpo = str(m._renderizar_corpo_login(None))
    for perfil in constantes.PERFIS_ACESSO:
        assert perfil in corpo


def test_renderizar_corpo_login_com_provisorio_mostra_profissionais():
    corpo = str(m._renderizar_corpo_login("Médico"))
    assert "Voltar" in corpo
    assert "Perfil Médico" in corpo


def test_voltar_login_limpa_o_provisorio():
    assert m._voltar_login(0) is dash.no_update
    assert m._voltar_login(1) is None


def test_barra_lateral_esconde_faturacao_para_enfermeiro():
    barra = str(m._barra_lateral("/", "Enfermeiro"))
    assert "Faturação" not in barra
    assert "Utentes" in barra


def test_barra_lateral_mostra_faturacao_para_recepcao_e_admin():
    barra_recepcao = str(m._barra_lateral("/", "Receção"))
    barra_admin = str(m._barra_lateral("/", "Admin"))
    assert "Faturação" in barra_recepcao
    assert "Faturação" in barra_admin


def test_barra_lateral_mostra_profissionais_so_para_admin():
    barra_admin = str(m._barra_lateral("/", "Admin"))
    barra_recepcao = str(m._barra_lateral("/", "Receção"))
    assert "Profissionais" in barra_admin
    assert "Profissionais" not in barra_recepcao


def test_rotear_pagina_sem_perfil_mostra_login():
    pagina = str(m._rotear_pagina("/utentes", None, None, None))
    assert "perfil" in pagina.lower() or "Enfermeiro" in pagina


def test_rotear_pagina_com_perfil_mostra_conteudo_pedido():
    pagina = str(m._rotear_pagina("/utentes", None, "Médico", None))
    assert "Nenhum utente encontrado" not in pagina  # é a página de utentes, não a de login
    assert "corpo-tabela-utentes" in pagina


def test_rotear_pagina_profissionais_recusa_recepcao_mas_permite_admin():
    pagina_recepcao = str(m._rotear_pagina("/profissionais", None, "Receção", None))
    assert "Acesso não disponível" in pagina_recepcao
    assert "corpo-tabela-profissionais" not in pagina_recepcao

    pagina_admin = str(m._rotear_pagina("/profissionais", None, "Admin", None))
    assert "corpo-tabela-profissionais" in pagina_admin


def test_rotear_pagina_inicio_mostra_banner_de_orientacao_uma_vez():
    pagina_primeira_vez = str(m._rotear_pagina("/", None, "Receção", None))
    assert "Entendi, não mostrar mais" in pagina_primeira_vez

    pagina_dispensada = str(m._rotear_pagina("/", None, "Receção", {"Receção": True}))
    assert "Entendi, não mostrar mais" not in pagina_dispensada

    # Dispensar num perfil não esconde a orientação de outro perfil.
    pagina_outro_perfil = str(m._rotear_pagina("/", None, "Admin", {"Receção": True}))
    assert "Entendi, não mostrar mais" in pagina_outro_perfil


def test_pagina_visao_geral_renderiza_com_kpis():
    pagina = str(m._pagina_visao_geral())
    assert "Visão Geral" in pagina


def test_pagina_visao_geral_sem_perfil_nao_mostra_banner():
    # Sessões antigas (sem perfil no _rotear_pagina) não devem partir.
    assert "banner-orientacao" not in str(m._pagina_visao_geral())


def test_banner_orientacao_tem_texto_para_os_quatro_perfis():
    for perfil in constantes.PERFIS_ACESSO:
        assert m._banner_orientacao(perfil) is not None


def test_dispensar_orientacao_grava_por_perfil_sem_apagar_outros():
    resultado = m._dispensar_orientacao(1, {"perfil": "Enfermeiro", "profissional": "Enf. Marta Sousa"}, {"Admin": True})
    assert resultado == {"Admin": True, "Enfermeiro": True}


def test_dispensar_orientacao_sem_cliques_nao_atualiza():
    assert m._dispensar_orientacao(0, {"perfil": "Enfermeiro", "profissional": None}, None) is dash.no_update


# --- Vitrine (landing page pública) e Simulação -------------------------------


def test_pagina_vitrine_tem_os_dois_botoes_de_entrada():
    pagina = str(m._pagina_vitrine())
    assert "href='/simulacao'" in pagina
    assert "href='/login'" in pagina
    assert "Simulação" in pagina
    assert "Ver o sistema inteiro" in pagina


def test_rotear_pagina_raiz_sem_sessao_mostra_vitrine():
    pagina = str(m._rotear_pagina("/", None, None, None))
    assert "Simulação" in pagina
    assert "href='/simulacao'" in pagina


def test_rotear_pagina_vitrine_e_publica_mesmo_sem_sessao():
    pagina = str(m._rotear_pagina("/vitrine", None, None, None))
    assert "Simulação" in pagina


def test_rotear_pagina_utentes_sem_sessao_continua_a_mostrar_login():
    # A raiz ("/") e o "/sobre" mostram a vitrine/sobre a visitantes sem
    # sessão — qualquer outro caminho continua a cair no login, tal como antes.
    pagina = str(m._rotear_pagina("/utentes", None, None, None))
    assert "perfil" in pagina.lower() or "Enfermeiro" in pagina


def test_rotear_pagina_sobre_e_publica_mesmo_sem_sessao():
    pagina = str(m._rotear_pagina("/sobre", None, None, None))
    assert "Sobre este projeto" in pagina
    assert "href='/vitrine'" in pagina


def test_pagina_sobre_logada_nao_mostra_ligacao_de_volta_a_vitrine():
    # Quem já está numa sessão vê a barra lateral normal — a ligação "Voltar
    # à vitrine" só faz sentido para quem chegou ao /sobre sem sessão.
    pagina = str(m._pagina_sobre(logado=True))
    assert "href='/vitrine'" not in pagina


def test_pagina_login_simulacao_tem_texto_proprio_e_ligacao_de_volta():
    pagina_normal = str(m._pagina_login())
    pagina_simulacao = str(m._pagina_login(simulacao=True))
    assert "Modo simulação" not in pagina_normal
    assert "Modo simulação" in pagina_simulacao
    assert "href='/vitrine'" in pagina_simulacao


def test_rotear_pagina_simulacao_mostra_escolha_de_perfil_em_modo_simulacao():
    pagina = str(m._rotear_pagina("/simulacao", None, None, None))
    assert "Modo simulação" in pagina
    # o corpo-login (escolha de perfil) só é preenchido pelo callback
    # _renderizar_corpo_login — a mesma função usada em /login — já coberto
    # por test_corpo_login_escolha_perfil_mostra_os_quatro_perfis.
    corpo = str(m._renderizar_corpo_login(None))
    for perfil in constantes.PERFIS_ACESSO:
        assert perfil in corpo


def test_e_simulacao_le_a_flag_gravada_na_sessao():
    assert m._e_simulacao({"perfil": "Receção", "profissional": None, "simulacao": True}) is True
    assert m._e_simulacao({"perfil": "Receção", "profissional": None, "simulacao": False}) is False
    assert m._e_simulacao({"perfil": "Receção", "profissional": None}) is False
    assert m._e_simulacao("Receção") is False  # sessões antigas (string simples)
    assert m._e_simulacao(None) is False


# _escolher_perfil_login e _escolher_profissional_login dependem de
# dash.ctx.triggered_id (preenchido pelo Dash a partir do componente que
# disparou o clique, algo que só existe num callback real em execução — tal
# como já acontecia antes desta funcionalidade, não são testados diretamente
# aqui) — a propagação da flag "simulacao" através deles está coberta pela
# verificação Playwright (clicar em "Simulação" na vitrine → escolher perfil
# → sessão gravada com simulacao=True → banner visível).


def test_sair_sessao_de_simulacao_volta_para_a_vitrine():
    _sessao, destino = m._sair_sessao(1, {"perfil": "Receção", "profissional": None, "simulacao": True})
    assert destino == "/vitrine"


def test_sair_sessao_normal_volta_para_o_login():
    _sessao, destino = m._sair_sessao(1, {"perfil": "Receção", "profissional": None, "simulacao": False})
    assert destino == "/login"


def test_rotear_pagina_mostra_banner_de_simulacao_fora_das_paginas_publicas():
    sessao = {"perfil": "Receção", "profissional": None, "simulacao": True}
    pagina_consultas = str(m._rotear_pagina("/consultas", None, sessao, None))
    assert "banner-simulacao" in pagina_consultas
    assert "Modo simulação" in pagina_consultas

    # nas páginas públicas o aviso não aparece — lá já é óbvio que ainda não
    # se entrou em nenhuma área da app
    pagina_vitrine = str(m._rotear_pagina("/vitrine", None, sessao, None))
    assert "banner-simulacao" not in pagina_vitrine


def test_rotear_pagina_sessao_normal_nunca_mostra_banner_de_simulacao():
    sessao = {"perfil": "Receção", "profissional": None, "simulacao": False}
    pagina = str(m._rotear_pagina("/consultas", None, sessao, None))
    assert "banner-simulacao" not in pagina


def test_barra_lateral_escondida_nas_paginas_publicas_mesmo_com_sessao():
    sessao = {"perfil": "Receção", "profissional": None, "simulacao": True}
    assert m._atualizar_barra_lateral("/vitrine", sessao) is None
    assert m._atualizar_barra_lateral("/simulacao", sessao) is None
    assert m._atualizar_barra_lateral("/login", sessao) is None
    assert m._atualizar_barra_lateral("/consultas", sessao) is not None


def test_ids_utentes_risco_elevado_e_subconjunto_dos_utentes():
    ids_risco = m._ids_utentes_risco_elevado()
    assert set(ids_risco) <= set(m.DADOS["utentes"]["id_utente"])


def test_pagina_sobre_renderiza():
    pagina = str(m._pagina_sobre())
    assert "Sobre este projeto" in pagina


def test_aba_historico_renderiza_para_utente_com_historico():
    utente = _primeiro_utente_com("historico")
    assert m._aba_historico(utente["id_utente"]) is not None


def test_gerar_pdf_ficha_produz_pdf_valido():
    utente = _primeiro_utente()
    conteudo = m._gerar_pdf_ficha(utente["id_utente"])
    assert conteudo[:4] == b"%PDF"


def test_vista_calendario_de_consultas_nao_rebenta():
    resultado = m._filtrar_consultas(None, None, None, "calendario", None)
    assert resultado is not None


def test_exportar_csv_faturacao_usa_ponto_virgula_e_virgula_decimal():
    # Excel em português usa ";" para separar colunas e "," como separador
    # decimal (ver app.py) — sem isto, ou o ficheiro abre tudo numa coluna
    # só, ou os números com casas decimais aparecem como texto.
    resultado = m._descarregar_csv_faturacao(1)
    linhas = resultado["content"].splitlines()
    cabecalho = linhas[0]
    primeira_linha_dados = linhas[1]
    assert ";" in cabecalho
    assert "," not in cabecalho  # nenhum título de coluna tem vírgula
    assert ";" in primeira_linha_dados
    assert "," in primeira_linha_dados  # algum valor com casas decimais usa vírgula
    assert "." not in primeira_linha_dados  # não deve sobrar nenhum ponto decimal


# --- Faturação: exportação em Excel -------------------------------------------


def test_exportar_excel_faturacao_produz_ficheiro_xlsx_valido():
    conteudo = m._gerar_excel_faturacao()
    # .xlsx é um zip — todo ficheiro zip começa por esta assinatura "PK".
    assert conteudo[:2] == b"PK"
    from openpyxl import load_workbook

    livro = load_workbook(BytesIO(conteudo))
    folha = livro.active
    cabecalho = [c.value for c in folha[1]]
    assert "Utente" in cabecalho
    assert folha.max_row == len(m.DADOS["faturacao"]) + 1  # +1 do cabeçalho


# --- Agendamento: criar, cancelar e reagendar consulta ------------------------


def _medico_ativo_para_teste():
    prof = m.DADOS["profissionais"]
    return prof[(prof["categoria"] == "Médico") & (prof["estado"] == "Ativo")].iloc[0]


def test_criar_consulta_sem_conflito_e_bem_sucedida():
    medico = _medico_ativo_para_teste()
    utente = _primeiro_utente()
    antes = len(m.DADOS["consultas"])
    mensagem, versao = m._criar_consulta(1, utente["id_utente"], medico["especialidade"], medico["nome"], "2027-01-10", "09:00", "Sala 1", 0)
    assert "risco_baixo" in mensagem.className
    assert len(m.DADOS["consultas"]) == antes + 1
    assert versao == 1


def test_criar_consulta_com_conflito_de_profissional_e_bloqueada():
    medico = _medico_ativo_para_teste()
    utente = _primeiro_utente()
    m._criar_consulta(1, utente["id_utente"], medico["especialidade"], medico["nome"], "2027-01-11", "09:00", "Sala 2", 0)
    antes = len(m.DADOS["consultas"])
    mensagem, versao = m._criar_consulta(1, utente["id_utente"], medico["especialidade"], medico["nome"], "2027-01-11", "09:00", "Sala 3", 0)
    assert "risco_elevado" in mensagem.className
    assert versao is dash.no_update
    assert len(m.DADOS["consultas"]) == antes  # não criou a segunda


def test_criar_consulta_com_campos_em_falta_pede_para_preencher():
    mensagem, versao = m._criar_consulta(1, None, None, None, None, None, None, 0)
    assert "risco_moderado" in mensagem.className
    assert versao is dash.no_update


def test_cancelar_consulta_muda_estado_para_cancelada():
    medico = _medico_ativo_para_teste()
    utente = _primeiro_utente()
    m._criar_consulta(1, utente["id_utente"], medico["especialidade"], medico["nome"], "2027-01-12", "10:00", "Sala 1", 0)
    id_consulta = m.DADOS["consultas"].iloc[-1]["id_consulta"]
    mensagem, _versao = m._cancelar_consulta(1, id_consulta, 0)
    assert "cancelada" in mensagem.children.lower()
    assert m.DADOS["consultas"].loc[m.DADOS["consultas"]["id_consulta"] == id_consulta, "estado"].iloc[0] == "Cancelada"


def test_reagendar_consulta_atualiza_data_hora_e_sala():
    medico = _medico_ativo_para_teste()
    utente = _primeiro_utente()
    m._criar_consulta(1, utente["id_utente"], medico["especialidade"], medico["nome"], "2027-01-13", "11:00", "Sala 1", 0)
    id_consulta = m.DADOS["consultas"].iloc[-1]["id_consulta"]
    mensagem, _versao = m._reagendar_consulta(1, id_consulta, "2027-01-14", "15:00", "Sala 4", 0)
    assert "risco_baixo" in mensagem.className
    linha = m.DADOS["consultas"].loc[m.DADOS["consultas"]["id_consulta"] == id_consulta].iloc[0]
    assert linha["sala"] == "Sala 4"
    assert pd.Timestamp(linha["data_hora"]) == pd.Timestamp("2027-01-14 15:00")


# --- Profissionais: CRUD -------------------------------------------------------


def test_profissionais_carregados_tem_pelo_menos_um_medico_ativo():
    prof = m.DADOS["profissionais"]
    assert ((prof["categoria"] == "Médico") & (prof["estado"] == "Ativo")).any()


def test_criar_profissional_medico_sem_especialidade_e_recusado():
    mensagem, _versao = m._criar_profissional(1, "Dr. Teste Sem Especialidade", "Médico", None, "x@x.com", "CP-1", "2020-01-01", 0)
    assert "risco_moderado" in mensagem.className


def test_criar_profissional_com_nome_duplicado_e_recusado():
    nome_existente = m.DADOS["profissionais"].iloc[0]["nome"]
    mensagem, _versao = m._criar_profissional(1, nome_existente, "Enfermeiro", None, "x@x.com", "CP-2", "2020-01-01", 0)
    assert "risco_elevado" in mensagem.className


def test_criar_e_remover_profissional_sem_historico():
    mensagem, versao = m._criar_profissional(1, "Enf. Teste Removível", "Enfermeiro", None, "teste@x.com", "CP-9", "2021-01-01", 0)
    assert "risco_baixo" in mensagem.className
    assert "Enf. Teste Removível" in m.DADOS["profissionais"]["nome"].values
    assert not m._profissional_tem_historico("Enf. Teste Removível")
    mensagem2, _versao2 = m._remover_profissional(1, "Enf. Teste Removível", versao)
    assert "risco_baixo" in mensagem2.className
    assert "Enf. Teste Removível" not in m.DADOS["profissionais"]["nome"].values


def test_remover_profissional_com_historico_e_bloqueado():
    nome_com_consultas = m.DADOS["consultas"].iloc[0]["profissional"]
    assert m._profissional_tem_historico(nome_com_consultas)
    mensagem, versao = m._remover_profissional(1, nome_com_consultas, 0)
    assert "Inativo" in mensagem.children
    assert versao is dash.no_update
    # continua no cadastro — não foi removido
    assert nome_com_consultas in m.DADOS["profissionais"]["nome"].values


def test_alternar_estado_profissional_liga_desliga():
    m._criar_profissional(1, "Fisio. Teste Estado", "Fisioterapeuta", None, "x@x.com", "CP-3", "2020-01-01", 0)
    estado_antes = m.DADOS["profissionais"].loc[m.DADOS["profissionais"]["nome"] == "Fisio. Teste Estado", "estado"].iloc[0]
    m._alternar_estado_profissional(1, "Fisio. Teste Estado", 0)
    estado_depois = m.DADOS["profissionais"].loc[m.DADOS["profissionais"]["nome"] == "Fisio. Teste Estado", "estado"].iloc[0]
    assert estado_antes != estado_depois


def test_editar_profissional_atualiza_contacto():
    m._criar_profissional(1, "Psic. Teste Editar", "Psicólogo", None, "antigo@x.com", "CP-4", "2020-01-01", 0)
    m._editar_profissional(1, "Psic. Teste Editar", "novo@x.com", None, None, 0)
    contacto = m.DADOS["profissionais"].loc[m.DADOS["profissionais"]["nome"] == "Psic. Teste Editar", "contacto"].iloc[0]
    assert contacto == "novo@x.com"


def test_pagina_profissionais_renderiza():
    assert m._pagina_profissionais() is not None


# --- Testes: acesso diferenciado por perfil (RBAC) --------------------------


def test_ficha_utente_esconde_avaliacao_risco_para_recepcao():
    id_utente = _primeiro_utente()["id_utente"]
    html_recepcao = str(m._pagina_ficha_utente(id_utente, "Receção"))
    assert "Avaliação de risco" not in html_recepcao


def test_ficha_utente_mostra_avaliacao_risco_para_medico_enfermeiro_e_admin():
    id_utente = _primeiro_utente()["id_utente"]
    assert "Avaliação de risco" in str(m._pagina_ficha_utente(id_utente, "Médico"))
    assert "Avaliação de risco" in str(m._pagina_ficha_utente(id_utente, "Enfermeiro"))
    assert "Avaliação de risco" in str(m._pagina_ficha_utente(id_utente, "Admin"))
    # Sem perfil definido (sessões antigas) mantém o comportamento anterior: mostra.
    assert "Avaliação de risco" in str(m._pagina_ficha_utente(id_utente))


def test_pagina_consultas_esconde_paineis_de_gestao_para_enfermeiro_e_medico():
    # Marcar/cancelar/reagendar é tarefa de agenda/front-desk — só a
    # Receção e o Admin. O Médico continua a ver a sua agenda (mais abaixo,
    # test_filtrar_consultas_medico_ve_apenas_a_sua_agenda), só não a gere.
    for perfil in ("Enfermeiro", "Médico"):
        pagina = str(m._pagina_consultas(perfil))
        assert "+ Nova Consulta" not in pagina
        assert "Gerir consulta existente" not in pagina


def test_pagina_consultas_mostra_paineis_de_gestao_para_recepcao_e_admin():
    html_recepcao = str(m._pagina_consultas("Receção"))
    html_admin = str(m._pagina_consultas("Admin"))
    for pagina in (html_recepcao, html_admin):
        assert "+ Nova Consulta" in pagina
        assert "Gerir consulta existente" in pagina


def test_filtrar_consultas_medico_ve_apenas_a_sua_agenda():
    nome_medico = m.DADOS["profissionais"].loc[
        (m.DADOS["profissionais"]["categoria"] == "Médico") & (m.DADOS["profissionais"]["estado"] == "Ativo"), "nome"
    ].iloc[0]
    total_do_medico = int((m.DADOS["consultas"]["profissional"] == nome_medico).sum())

    resultado = m._filtrar_consultas(None, None, None, sessao={"perfil": "Médico", "profissional": nome_medico})
    resultado_texto = str(resultado)

    outros_medicos = m.DADOS["profissionais"].loc[
        (m.DADOS["profissionais"]["categoria"] == "Médico") & (m.DADOS["profissionais"]["nome"] != nome_medico), "nome"
    ]
    if total_do_medico and len(outros_medicos):
        assert nome_medico in resultado_texto
        # Nenhum outro médico deve aparecer na tabela filtrada.
        assert not any(outro in resultado_texto for outro in outros_medicos)


def test_filtrar_consultas_sem_perfil_medico_mostra_todas():
    resultado_recepcao = str(m._filtrar_consultas(None, None, None, sessao={"perfil": "Receção", "profissional": None}))
    resultado_sem_sessao = str(m._filtrar_consultas(None, None, None))
    # Sem filtragem por agenda própria, ambos devolvem a mesma tabela completa.
    assert resultado_recepcao == resultado_sem_sessao


def test_opcoes_consultas_para_gerir_filtra_por_medico_da_sessao():
    nome_medico = m.DADOS["profissionais"].loc[
        (m.DADOS["profissionais"]["categoria"] == "Médico") & (m.DADOS["profissionais"]["estado"] == "Ativo"), "nome"
    ].iloc[0]
    opcoes_medico = m._opcoes_consultas_para_gerir(0, sessao={"perfil": "Médico", "profissional": nome_medico})
    opcoes_todas = m._opcoes_consultas_para_gerir(0, sessao={"perfil": "Receção", "profissional": None})
    assert len(opcoes_medico) <= len(opcoes_todas)


# --- Testes: prescrições e receita em PDF (só Médico) ------------------------


def test_aba_e_corpo_prescricoes_esconde_formulario_e_receita_para_nao_medico():
    utente = _primeiro_utente()
    for perfil, profissional in [("Enfermeiro", "Enf. Marta Sousa"), ("Receção", None), ("Admin", None)]:
        aba = str(m._aba_prescricoes(utente["id_utente"], perfil, profissional))
        corpo = str(m._corpo_prescricoes(utente["id_utente"], perfil, profissional))
        assert "+ Nova Prescrição" not in aba
        assert "Gerar receita em PDF" not in corpo


def test_aba_e_corpo_prescricoes_mostra_formulario_e_receita_para_medico():
    utente = _primeiro_utente()
    medico = _medico_ativo_para_teste()
    aba = str(m._aba_prescricoes(utente["id_utente"], "Médico", medico["nome"]))
    corpo = str(m._corpo_prescricoes(utente["id_utente"], "Médico", medico["nome"]))
    assert "+ Nova Prescrição" in aba
    assert "Gerar receita em PDF" in corpo


def test_criar_prescricao_como_medico_e_bem_sucedida():
    medico = _medico_ativo_para_teste()
    utente = _primeiro_utente()
    antes = len(m.DADOS["prescricoes"])
    mensagem, versao = m._criar_prescricao(
        1, utente["id_utente"], "Ibuprofeno", "400mg, 2x/dia", "5 dias",
        {"perfil": "Médico", "profissional": medico["nome"]}, 0,
    )
    assert "risco_baixo" in mensagem.className
    assert len(m.DADOS["prescricoes"]) == antes + 1
    assert versao == 1
    nova = m.DADOS["prescricoes"].iloc[-1]
    assert nova["medicamento"] == "Ibuprofeno"
    assert nova["profissional"] == medico["nome"]


def test_criar_prescricao_fora_do_perfil_medico_e_recusada():
    utente = _primeiro_utente()
    antes = len(m.DADOS["prescricoes"])
    mensagem, versao = m._criar_prescricao(
        1, utente["id_utente"], "Paracetamol", "500mg, 1x/dia", None,
        {"perfil": "Receção", "profissional": None}, 0,
    )
    assert "risco_elevado" in mensagem.className
    assert versao is dash.no_update
    assert len(m.DADOS["prescricoes"]) == antes


def test_criar_prescricao_com_campos_em_falta_pede_para_preencher():
    medico = _medico_ativo_para_teste()
    utente = _primeiro_utente()
    mensagem, versao = m._criar_prescricao(
        1, utente["id_utente"], None, None, None,
        {"perfil": "Médico", "profissional": medico["nome"]}, 0,
    )
    assert "risco_moderado" in mensagem.className
    assert versao is dash.no_update


def test_gerar_pdf_receita_produz_pdf_valido():
    medico = _medico_ativo_para_teste()
    utente = _primeiro_utente()
    m._criar_prescricao(
        1, utente["id_utente"], "Amoxicilina", "500mg, 3x/dia", "7 dias",
        {"perfil": "Médico", "profissional": medico["nome"]}, 0,
    )
    id_prescricao = m.DADOS["prescricoes"].iloc[-1]["id_prescricao"]
    pdf_bytes = m._gerar_pdf_receita(id_prescricao)
    assert pdf_bytes is not None
    assert pdf_bytes[:4] == b"%PDF"


def test_gerar_pdf_receita_com_id_inexistente_devolve_none():
    assert m._gerar_pdf_receita("RX99999") is None


# --- Testes: relatórios (atividade profissional + consulta) ------------------


def test_estatisticas_atividade_profissional_calcula_totais_e_taxa():
    medico = _medico_ativo_para_teste()
    utente = _primeiro_utente()
    m._criar_consulta(1, utente["id_utente"], medico["especialidade"], medico["nome"], "2027-02-01", "09:00", "Sala 1", 0)
    id_consulta = m.DADOS["consultas"].iloc[-1]["id_consulta"]
    m.DADOS["consultas"].loc[m.DADOS["consultas"]["id_consulta"] == id_consulta, "estado"] = "Realizada"

    stats = m._estatisticas_atividade_profissional(medico["nome"])
    assert stats["total"] >= 1
    assert stats["por_estado"].get("Realizada", 0) >= 1
    assert stats["taxa_comparencia"] is not None


def test_estatisticas_atividade_profissional_sem_consultas_no_periodo():
    stats = m._estatisticas_atividade_profissional("Nome Que Não Existe")
    assert stats["total"] == 0
    assert stats["taxa_comparencia"] is None
    assert stats["por_especialidade"] == {}


def test_opcoes_relatorio_consulta_filtra_por_medico_da_sessao():
    nome_medico = m.DADOS["profissionais"].loc[
        (m.DADOS["profissionais"]["categoria"] == "Médico") & (m.DADOS["profissionais"]["estado"] == "Ativo"), "nome"
    ].iloc[0]
    opcoes_medico = m._opcoes_relatorio_consulta("Médico", nome_medico)
    opcoes_todas = m._opcoes_relatorio_consulta("Receção", None)
    assert len(opcoes_medico) <= len(opcoes_todas)


def test_pagina_relatorios_admin_permite_escolher_profissional():
    pagina = str(m._pagina_relatorios("Admin", None, None))
    assert "select-relatorio-profissional" in pagina


def test_pagina_relatorios_le_profissional_da_query_string():
    medico = _medico_ativo_para_teste()
    pagina = m._pagina_relatorios("Admin", None, f"?profissional={medico['nome']}")
    assert medico["nome"] in str(pagina)


def test_pagina_relatorios_recepcao_nao_tem_atividade_profissional():
    # A Receção só vê o relatório de consulta — a atividade agregada de um
    # profissional é informação de gestão, reservada ao Admin.
    pagina = str(m._pagina_relatorios("Receção", None, None))
    assert "select-relatorio-profissional" not in pagina
    assert "Relatório de atividade profissional" not in pagina
    assert "select-relatorio-consulta" in pagina


def test_pagina_relatorios_medico_fica_preso_a_propria_identidade():
    medico = _medico_ativo_para_teste()
    # Mesmo que a query string tente indicar outro profissional, o Médico
    # só pode ver o relatório da sua própria identidade autenticada.
    pagina = str(m._pagina_relatorios("Médico", medico["nome"], "?profissional=Outra Pessoa"))
    assert "Outra Pessoa" not in pagina


def test_gerar_pdf_relatorio_atividade_produz_pdf_valido():
    medico = _medico_ativo_para_teste()
    pdf_bytes = m._gerar_pdf_relatorio_atividade(medico["nome"])
    assert pdf_bytes[:4] == b"%PDF"


def test_gerar_excel_relatorio_atividade_produz_xlsx_valido():
    from openpyxl import load_workbook

    medico = _medico_ativo_para_teste()
    conteudo = m._gerar_excel_relatorio_atividade(medico["nome"])
    livro = load_workbook(BytesIO(conteudo))
    assert "Resumo" in livro.sheetnames
    assert "Por especialidade" in livro.sheetnames


def test_gerar_pdf_relatorio_consulta_produz_pdf_valido():
    medico = _medico_ativo_para_teste()
    utente = _primeiro_utente()
    m._criar_consulta(1, utente["id_utente"], medico["especialidade"], medico["nome"], "2027-02-02", "10:00", "Sala 2", 0)
    id_consulta = m.DADOS["consultas"].iloc[-1]["id_consulta"]
    pdf_bytes = m._gerar_pdf_relatorio_consulta(id_consulta)
    assert pdf_bytes is not None
    assert pdf_bytes[:4] == b"%PDF"


def test_gerar_pdf_relatorio_consulta_com_id_inexistente_devolve_none():
    assert m._gerar_pdf_relatorio_consulta("C99999") is None


def test_gerar_excel_relatorio_consulta_produz_xlsx_valido():
    from openpyxl import load_workbook

    medico = _medico_ativo_para_teste()
    utente = _primeiro_utente()
    m._criar_consulta(1, utente["id_utente"], medico["especialidade"], medico["nome"], "2027-02-03", "11:00", "Sala 3", 0)
    id_consulta = m.DADOS["consultas"].iloc[-1]["id_consulta"]
    conteudo = m._gerar_excel_relatorio_consulta(id_consulta)
    livro = load_workbook(BytesIO(conteudo))
    assert livro.active.title == "Consulta"


def test_gerar_excel_relatorio_consulta_com_id_inexistente_devolve_none():
    assert m._gerar_excel_relatorio_consulta("C99999") is None
