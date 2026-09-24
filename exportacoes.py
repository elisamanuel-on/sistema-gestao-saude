"""Geração dos documentos exportáveis (PDF em reportlab, Excel .xlsx em
openpyxl): ficha clínica, receita médica, relatório de atividade
profissional e relatório de consulta, faturação. Cada função devolve os
bytes prontos a enviar (dcc.send_bytes, nos callbacks de download)."""

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas as pdf_canvas

from nucleo import COLUNAS_EXPORTACAO_FATURACAO, DADOS, _consultas_periodo_profissional, _estatisticas_atividade_profissional, _idade_utente

def _gerar_excel_faturacao():
    """Excel (.xlsx) a sério, não um CSV disfarçado — resolve de vez a
    ambiguidade da vírgula/ponto e vírgula (essa depende só da configuração
    regional do computador de quem abre o CSV, não do ficheiro em si), e de
    caminho já formata cabeçalhos e moeda."""
    fat = DADOS["faturacao"].merge(DADOS["utentes"][["id_utente", "nome", "tipo_cuidado"]], on="id_utente")
    fat = fat[list(COLUNAS_EXPORTACAO_FATURACAO.keys())].rename(columns=COLUNAS_EXPORTACAO_FATURACAO)
    colunas_euro = {c for c in fat.columns if "€" in c}

    livro = Workbook()
    folha = livro.active
    folha.title = "Faturação"
    folha.append(list(fat.columns))
    for celula in folha[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="0F7A6C")
    for _idx, linha in fat.iterrows():
        folha.append(list(linha))

    for i, coluna in enumerate(fat.columns, start=1):
        letra = get_column_letter(i)
        largura = max(len(coluna), int(fat[coluna].astype(str).map(len).max())) + 3
        folha.column_dimensions[letra].width = min(34, max(11, largura))
        if coluna in colunas_euro:
            for linha_celulas in folha.iter_rows(min_row=2, min_col=i, max_col=i):
                linha_celulas[0].number_format = "#,##0.00 €"
    folha.freeze_panes = "A2"

    buffer = io.BytesIO()
    livro.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _escrever_secoes_clinicas_utente(escrever, id_utente):
    """Blocos de alergias / diagnósticos / plano de cuidados / últimos sinais
    vitais de um utente, escritos num PDF através da função `escrever` do
    documento que estiver a ser gerado. Partilhado entre a ficha clínica e o
    relatório de consulta, para os dois mostrarem sempre a mesma informação
    clínica de um utente (em vez de duas versões que podiam divergir)."""
    alergias = DADOS["alergias"][DADOS["alergias"]["id_utente"] == id_utente]["alergia"].tolist()
    escrever("Alergias", 12, negrito=True)
    escrever(", ".join(alergias) if alergias else "Nenhuma conhecida", 10, espaco=10 * mm)

    diagnosticos = DADOS["diagnosticos"][DADOS["diagnosticos"]["id_utente"] == id_utente].sort_values("data", ascending=False)
    escrever("Diagnósticos", 12, negrito=True)
    if diagnosticos.empty:
        escrever("Sem diagnósticos registados.", 10)
    for _idx, d in diagnosticos.iterrows():
        escrever(f"- {d['data']}: {d['diagnostico']} ({d['profissional']})", 9, espaco=5 * mm)

    plano = DADOS["plano_cuidados"][DADOS["plano_cuidados"]["id_utente"] == id_utente]
    escrever("Plano de cuidados multidisciplinar", 12, negrito=True)
    if plano.empty:
        escrever("Sem plano de cuidados registado.", 10)
    for _idx, p in plano.iterrows():
        escrever(
            f"- [{p['area_profissional']}] {p['objetivo']}: {p['estado']} ({p['progresso_percent']}%) · "
            f"Resp.: {p['profissional_responsavel']}",
            9,
            espaco=5 * mm,
        )

    ultimos_vitais = DADOS["vitais"][DADOS["vitais"]["id_utente"] == id_utente].sort_values("data_hora").tail(1)
    escrever("Últimos sinais vitais", 12, negrito=True)
    if len(ultimos_vitais):
        v = ultimos_vitais.iloc[0]
        escrever(
            f"Glicemia: {v['glicemia']} mg/dL · Peso: {v['peso_kg']} kg · "
            f"Tensão: {v['tensao_sistolica']}/{v['tensao_diastolica']} mmHg · Temp.: {v['temperatura']}°C",
            10,
        )
    else:
        escrever("Sem sinais vitais registados.", 10)


def _gerar_pdf_ficha(id_utente):
    utente = DADOS["utentes"].loc[DADOS["utentes"]["id_utente"] == id_utente].iloc[0]
    buffer = io.BytesIO()
    c = pdf_canvas.Canvas(buffer, pagesize=A4)
    largura, altura = A4
    margem_esquerda = 20 * mm
    posicao_y = altura - 20 * mm

    def escrever(texto, tamanho=10, negrito=False, espaco=6 * mm):
        nonlocal posicao_y
        if posicao_y < 25 * mm:
            c.showPage()
            posicao_y = altura - 20 * mm
        c.setFont("Helvetica-Bold" if negrito else "Helvetica", tamanho)
        c.drawString(margem_esquerda, posicao_y, texto)
        posicao_y -= espaco

    escrever(f"Ficha clínica: {utente['nome']}", 16, negrito=True, espaco=10 * mm)
    escrever(
        f"{utente['genero']} · {_idade_utente(utente['data_nascimento'])} anos · "
        f"{utente['tipo_cuidado']} · Processo {utente['processo']}",
        10,
    )
    escrever(f"Estado: {utente['estado']}   Admissão: {utente['data_admissao']}", 10, espaco=10 * mm)

    _escrever_secoes_clinicas_utente(escrever, id_utente)

    c.setFont("Helvetica-Oblique", 8)
    c.drawString(
        margem_esquerda, 15 * mm,
        "Documento gerado automaticamente: demonstração técnica de portfólio, dados sintéticos, não usar com pacientes reais.",
    )

    c.save()
    buffer.seek(0)
    return buffer.getvalue()


def _gerar_pdf_receita(id_prescricao):
    """Gera uma receita médica em PDF a partir de uma prescrição registada.
    Só é chamada para prescrições que pertencem ao médico autenticado (ver
    filtro em _corpo_prescricoes / _descarregar_receita_pdf)."""
    linha = DADOS["prescricoes"].loc[DADOS["prescricoes"]["id_prescricao"] == id_prescricao]
    if linha.empty:
        return None
    p = linha.iloc[0]
    utente = DADOS["utentes"].loc[DADOS["utentes"]["id_utente"] == p["id_utente"]].iloc[0]
    prof = DADOS["profissionais"].loc[DADOS["profissionais"]["nome"] == p["profissional"]]
    especialidade = prof.iloc[0]["especialidade"] if len(prof) and prof.iloc[0]["especialidade"] else "Medicina Geral"
    numero_cedula = prof.iloc[0]["numero_cedula"] if len(prof) else "N/D"

    buffer = io.BytesIO()
    c = pdf_canvas.Canvas(buffer, pagesize=A4)
    largura, altura = A4
    margem_esquerda = 20 * mm
    posicao_y = altura - 20 * mm

    def escrever(texto, tamanho=10, negrito=False, espaco=6 * mm, centrado=False):
        nonlocal posicao_y
        c.setFont("Helvetica-Bold" if negrito else "Helvetica", tamanho)
        if centrado:
            c.drawCentredString(largura / 2, posicao_y, texto)
        else:
            c.drawString(margem_esquerda, posicao_y, texto)
        posicao_y -= espaco

    escrever("Receita Médica", 18, negrito=True, espaco=12 * mm, centrado=True)
    escrever(f"{p['profissional']} · {especialidade}", 11, negrito=True)
    escrever(f"Cédula profissional: {numero_cedula}", 9, espaco=10 * mm)

    escrever(f"Utente: {utente['nome']}", 11, negrito=True)
    escrever(
        f"{utente['genero']} · {_idade_utente(utente['data_nascimento'])} anos · Processo {utente['processo']}",
        9,
        espaco=12 * mm,
    )

    escrever("Prescrição", 12, negrito=True)
    escrever(p["medicamento"], 11, espaco=6 * mm)
    escrever(f"Posologia: {p['posologia']}", 10)
    if p.get("duracao"):
        escrever(f"Duração do tratamento: {p['duracao']}", 10)
    posicao_y -= 14 * mm

    escrever(f"Data de emissão: {p['data']}", 10, espaco=18 * mm)
    c.line(margem_esquerda, posicao_y, margem_esquerda + 70 * mm, posicao_y)
    posicao_y -= 5 * mm
    escrever("Assinatura e carimbo", 8)

    c.setFont("Helvetica-Oblique", 8)
    c.drawString(
        margem_esquerda, 15 * mm,
        "Documento gerado automaticamente: demonstração técnica de portfólio, dados sintéticos, não usar com pacientes reais.",
    )

    c.save()
    buffer.seek(0)
    return buffer.getvalue()


def _gerar_pdf_relatorio_atividade(nome, data_inicio=None, data_fim=None):
    stats = _estatisticas_atividade_profissional(nome, data_inicio, data_fim)
    prof = DADOS["profissionais"].loc[DADOS["profissionais"]["nome"] == nome]
    categoria = prof.iloc[0]["categoria"] if len(prof) else ""

    buffer = io.BytesIO()
    c = pdf_canvas.Canvas(buffer, pagesize=A4)
    largura, altura = A4
    margem_esquerda = 20 * mm
    posicao_y = altura - 20 * mm

    def escrever(texto, tamanho=10, negrito=False, espaco=6 * mm):
        nonlocal posicao_y
        if posicao_y < 25 * mm:
            c.showPage()
            posicao_y = altura - 20 * mm
        c.setFont("Helvetica-Bold" if negrito else "Helvetica", tamanho)
        c.drawString(margem_esquerda, posicao_y, texto)
        posicao_y -= espaco

    escrever("Relatório de Atividade Profissional", 16, negrito=True, espaco=10 * mm)
    escrever(f"{nome}" + (f" · {categoria}" if categoria else ""), 12, negrito=True)
    escrever(f"Período: {data_inicio or 'início dos registos'} a {data_fim or 'hoje'}", 10, espaco=10 * mm)

    escrever("Resumo", 12, negrito=True)
    escrever(f"Total de consultas: {stats['total']}", 10)
    for estado_nome, qtd in stats["por_estado"].items():
        escrever(f"- {estado_nome}: {qtd}", 9, espaco=5 * mm)
    taxa_texto = f"{stats['taxa_comparencia']:.0f}%" if stats["taxa_comparencia"] is not None else "sem dados suficientes"
    escrever(f"Taxa de comparência: {taxa_texto}", 10, espaco=10 * mm)

    escrever("Consultas por especialidade", 12, negrito=True)
    especialidades_ordenadas = sorted(stats["por_especialidade"].items(), key=lambda item: -item[1])
    if especialidades_ordenadas:
        for especialidade, qtd in especialidades_ordenadas:
            percent = (qtd / stats["total"] * 100) if stats["total"] else 0
            escrever(f"- {especialidade}: {qtd} ({percent:.0f}%)", 9, espaco=5 * mm)
    else:
        escrever("Sem consultas registadas no período.", 10)
    posicao_y -= 4 * mm

    escrever("Consultas do período (detalhe)", 12, negrito=True)
    detalhe = _consultas_periodo_profissional(nome, data_inicio, data_fim)
    if detalhe.empty:
        escrever("Sem consultas registadas no período.", 10)
    else:
        for _idx, r in detalhe.iterrows():
            escrever(
                f"- {r['data_hora'].strftime('%d/%m/%Y %H:%M')} · {r['nome']} · {r['especialidade']} · "
                f"{r['sala']} · {r['estado']}",
                9,
                espaco=5 * mm,
            )

    c.setFont("Helvetica-Oblique", 8)
    c.drawString(
        margem_esquerda, 15 * mm,
        "Documento gerado automaticamente: demonstração técnica de portfólio, dados sintéticos, não usar com pacientes reais.",
    )
    c.save()
    buffer.seek(0)
    return buffer.getvalue()


def _gerar_excel_relatorio_atividade(nome, data_inicio=None, data_fim=None):
    stats = _estatisticas_atividade_profissional(nome, data_inicio, data_fim)

    livro = Workbook()
    folha_resumo = livro.active
    folha_resumo.title = "Resumo"
    folha_resumo.append(["Métrica", "Valor"])
    for celula in folha_resumo[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="0F7A6C")
    folha_resumo.append(["Profissional", nome])
    folha_resumo.append(["Período", f"{data_inicio or 'início dos registos'} a {data_fim or 'hoje'}"])
    folha_resumo.append(["Total de consultas", stats["total"]])
    for estado_nome, qtd in stats["por_estado"].items():
        folha_resumo.append([estado_nome, qtd])
    taxa_texto = f"{stats['taxa_comparencia']:.0f}%" if stats["taxa_comparencia"] is not None else "N/D"
    folha_resumo.append(["Taxa de comparência", taxa_texto])
    folha_resumo.column_dimensions["A"].width = 24
    folha_resumo.column_dimensions["B"].width = 30

    folha_especialidade = livro.create_sheet("Por especialidade")
    folha_especialidade.append(["Especialidade", "Consultas"])
    for celula in folha_especialidade[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="0F7A6C")
    for especialidade, qtd in sorted(stats["por_especialidade"].items(), key=lambda item: -item[1]):
        folha_especialidade.append([especialidade, qtd])
    folha_especialidade.column_dimensions["A"].width = 22
    folha_especialidade.column_dimensions["B"].width = 14

    folha_detalhe = livro.create_sheet("Consultas do período")
    folha_detalhe.append(["Data/Hora", "Utente", "Especialidade", "Sala", "Estado"])
    for celula in folha_detalhe[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="0F7A6C")
    detalhe = _consultas_periodo_profissional(nome, data_inicio, data_fim)
    for _idx, r in detalhe.iterrows():
        folha_detalhe.append([r["data_hora"].strftime("%d/%m/%Y %H:%M"), r["nome"], r["especialidade"], r["sala"], r["estado"]])
    folha_detalhe.column_dimensions["A"].width = 18
    folha_detalhe.column_dimensions["B"].width = 28
    folha_detalhe.column_dimensions["C"].width = 18
    folha_detalhe.column_dimensions["D"].width = 10
    folha_detalhe.column_dimensions["E"].width = 14

    buffer = io.BytesIO()
    livro.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _gerar_pdf_relatorio_consulta(id_consulta):
    linha = DADOS["consultas"].loc[DADOS["consultas"]["id_consulta"] == id_consulta]
    if linha.empty:
        return None
    consulta = linha.iloc[0]
    utente = DADOS["utentes"].loc[DADOS["utentes"]["id_utente"] == consulta["id_utente"]].iloc[0]

    buffer = io.BytesIO()
    c = pdf_canvas.Canvas(buffer, pagesize=A4)
    _largura, altura = A4
    margem_esquerda = 20 * mm
    posicao_y = altura - 20 * mm

    def escrever(texto, tamanho=10, negrito=False, espaco=6 * mm):
        nonlocal posicao_y
        if posicao_y < 25 * mm:
            c.showPage()
            posicao_y = altura - 20 * mm
        c.setFont("Helvetica-Bold" if negrito else "Helvetica", tamanho)
        c.drawString(margem_esquerda, posicao_y, texto)
        posicao_y -= espaco

    escrever("Relatório de Consulta", 16, negrito=True, espaco=10 * mm)
    escrever(f"Utente: {utente['nome']}", 11, negrito=True)
    escrever(
        f"{utente['genero']} · {_idade_utente(utente['data_nascimento'])} anos · Processo {utente['processo']} · "
        f"{utente['tipo_cuidado']}",
        9,
    )
    escrever(f"Estado atual do utente: {utente['estado']}   Quarto: {utente['quarto']}", 9, espaco=10 * mm)

    escrever("Dados da consulta", 12, negrito=True)
    escrever(f"Data/Hora: {consulta['data_hora'].strftime('%d/%m/%Y %H:%M')}", 10)
    escrever(f"Especialidade: {consulta['especialidade']}", 10)
    escrever(f"Profissional: {consulta['profissional']}", 10)
    escrever(f"Sala: {consulta['sala']}", 10)
    escrever(f"Estado da consulta: {consulta['estado']}", 10, espaco=10 * mm)

    escrever("Informação clínica do utente", 12, negrito=True, espaco=8 * mm)
    _escrever_secoes_clinicas_utente(escrever, consulta["id_utente"])

    c.setFont("Helvetica-Oblique", 8)
    c.drawString(
        margem_esquerda, 15 * mm,
        "Documento gerado automaticamente: demonstração técnica de portfólio, dados sintéticos, não usar com pacientes reais.",
    )
    c.save()
    buffer.seek(0)
    return buffer.getvalue()


def _gerar_excel_relatorio_consulta(id_consulta):
    linha = DADOS["consultas"].loc[DADOS["consultas"]["id_consulta"] == id_consulta]
    if linha.empty:
        return None
    consulta = linha.iloc[0]
    utente = DADOS["utentes"].loc[DADOS["utentes"]["id_utente"] == consulta["id_utente"]].iloc[0]

    id_utente = consulta["id_utente"]

    livro = Workbook()
    folha = livro.active
    folha.title = "Consulta"
    folha.append(["Campo", "Valor"])
    for celula in folha[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="0F7A6C")
    folha.append(["Utente", utente["nome"]])
    folha.append(["Género", utente["genero"]])
    folha.append(["Idade", _idade_utente(utente["data_nascimento"])])
    folha.append(["Processo", utente["processo"]])
    folha.append(["Tipo de cuidado", utente["tipo_cuidado"]])
    folha.append(["Estado atual do utente", utente["estado"]])
    folha.append(["Quarto", utente["quarto"]])
    folha.append(["Data/Hora da consulta", consulta["data_hora"].strftime("%d/%m/%Y %H:%M")])
    folha.append(["Especialidade", consulta["especialidade"]])
    folha.append(["Profissional", consulta["profissional"]])
    folha.append(["Sala", consulta["sala"]])
    folha.append(["Estado da consulta", consulta["estado"]])
    folha.column_dimensions["A"].width = 22
    folha.column_dimensions["B"].width = 30

    def _folha_lista(titulo, cabecalhos, linhas, larguras):
        folha_nova = livro.create_sheet(titulo)
        folha_nova.append(cabecalhos)
        for celula in folha_nova[1]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor="0F7A6C")
        for linha_dados in linhas:
            folha_nova.append(linha_dados)
        for coluna, largura_coluna in zip("ABCDEFG", larguras, strict=False):
            folha_nova.column_dimensions[coluna].width = largura_coluna
        return folha_nova

    alergias = DADOS["alergias"][DADOS["alergias"]["id_utente"] == id_utente]["alergia"].tolist()
    _folha_lista("Alergias", ["Alergia"], [[a] for a in alergias], [30])

    diagnosticos = DADOS["diagnosticos"][DADOS["diagnosticos"]["id_utente"] == id_utente].sort_values("data", ascending=False)
    _folha_lista(
        "Diagnósticos",
        ["Data", "Diagnóstico", "Profissional"],
        [[d["data"], d["diagnostico"], d["profissional"]] for _idx, d in diagnosticos.iterrows()],
        [14, 32, 26],
    )

    plano = DADOS["plano_cuidados"][DADOS["plano_cuidados"]["id_utente"] == id_utente]
    _folha_lista(
        "Plano de cuidados",
        ["Área", "Objetivo", "Estado", "Progresso (%)", "Responsável"],
        [
            [p["area_profissional"], p["objetivo"], p["estado"], p["progresso_percent"], p["profissional_responsavel"]]
            for _idx, p in plano.iterrows()
        ],
        [16, 32, 16, 14, 26],
    )

    vitais = DADOS["vitais"][DADOS["vitais"]["id_utente"] == id_utente].sort_values("data_hora", ascending=False).head(5)
    _folha_lista(
        "Sinais vitais",
        ["Data/Hora", "Glicemia (mg/dL)", "Peso (kg)", "Tensão sistólica", "Tensão diastólica", "Temperatura (°C)"],
        [
            [
                v["data_hora"].strftime("%d/%m/%Y %H:%M") if hasattr(v["data_hora"], "strftime") else v["data_hora"],
                v["glicemia"],
                v["peso_kg"],
                v["tensao_sistolica"],
                v["tensao_diastolica"],
                v["temperatura"],
            ]
            for _idx, v in vitais.iterrows()
        ],
        [18, 16, 12, 16, 16, 16],
    )

    buffer = io.BytesIO()
    livro.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()



__all__ = [
    "_gerar_excel_faturacao",
    "_escrever_secoes_clinicas_utente",
    "_gerar_pdf_ficha",
    "_gerar_pdf_receita",
    "_gerar_pdf_relatorio_atividade",
    "_gerar_excel_relatorio_atividade",
    "_gerar_pdf_relatorio_consulta",
    "_gerar_excel_relatorio_consulta",
]
